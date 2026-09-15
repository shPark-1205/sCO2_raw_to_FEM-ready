import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from main import _macro_reference, _validate_range_integrity, _values_equivalent
from main import build_output_xlsm_path, find_actual_data_bounds
from main import parse_merged_filename, read_merged_values


SAMPLE_NAME = "20260914_7.771MPa_31.3C_40000_8V_Hd1.00_Merged.xlsx"


class FilenameTests(unittest.TestCase):
    def test_sample_filename(self):
        parsed = parse_merged_filename(SAMPLE_NAME)
        self.assertEqual(parsed["output_filename"], "20260914_7.771MPa_31.3C_40000_8V_Hd1.00.xlsm")
        self.assertEqual(parsed["date_text"], "26.09.14")
        self.assertEqual(parsed["voltage"], "8")
        self.assertEqual(parsed["hd"], "1.00")
        self.assertEqual(parsed["experiment_folder"], "20260914_7.771MPa_31.3C_40000")

    def test_decimal_voltage_and_variable_hd_precision(self):
        parsed = parse_merged_filename("20260914_7.771MPa_31.3C_40000_8.25V_Hd0.875_Merged.xlsx")
        self.assertEqual(parsed["voltage"], "8.25")
        self.assertEqual(parsed["hd"], "0.875")

    def test_invalid_filename_is_rejected(self):
        with self.assertRaises(ValueError):
            parse_merged_filename("unexpected_Merged.xlsx")

    def test_output_path(self):
        result = build_output_xlsm_path(SAMPLE_NAME, {"xlsm_output_dir": r"C:\results"})
        self.assertEqual(result.name, "20260914_7.771MPa_31.3C_40000_8V_Hd1.00.xlsm")
        self.assertEqual(result.parent.name, "20260914_7.771MPa_31.3C_40000")

    def test_macro_reference_escapes_apostrophe(self):
        self.assertEqual(_macro_reference("a'b.xlsm", "calculator"), "'a''b.xlsm'!calculator")


class WorkbookValueTests(unittest.TestCase):
    def test_excel_time_serial_matches_source_time_text(self):
        self.assertTrue(_values_equivalent("20:14:10", 0.8431712962962963))

    def test_range_integrity_rejects_partial_representative_match(self):
        with self.assertRaises(RuntimeError):
            _validate_range_integrity(100, 100, 1, 5, 5, "검증 실패")

    def test_range_integrity_rejects_large_counta_difference(self):
        with self.assertRaises(RuntimeError):
            _validate_range_integrity(10_000, 9_000, 5, 5, 5, "검증 실패")

    def test_actual_bounds_ignore_styled_empty_cell(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet["A1"] = 1
        worksheet["C2"] = 2
        worksheet["Z100"].number_format = "0.00"
        self.assertEqual(find_actual_data_bounds(worksheet), (2, 3))

    def test_read_values_preserves_internal_blanks(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "source.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "MergedData"
            worksheet["A1"] = "start"
            worksheet["C2"] = 7
            workbook.save(path)
            workbook.close()
            self.assertEqual(read_merged_values(path), (("start", None, None), (None, None, 7)))

    def test_empty_merged_data_is_rejected(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "empty.xlsx"
            workbook = Workbook()
            workbook.active.title = "MergedData"
            workbook.save(path)
            workbook.close()
            with self.assertRaises(ValueError):
                read_merged_values(path)


if __name__ == "__main__":
    unittest.main()
