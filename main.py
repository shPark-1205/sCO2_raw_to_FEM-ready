"""실험 원본 데이터를 병합하고 Excel 매크로 템플릿으로 후처리합니다.

처리 흐름은 크게 두 단계입니다.
1. RTD/SYS/HX 원본을 기존 규칙 그대로 결합해 ``*_Merged.xlsx``를 만듭니다.
2. 생성에 성공한 파일을 하나씩 XLSM 템플릿에 넣고 VBA 매크로를 실행합니다.

두 번째 단계는 Excel COM과 VBA 팝업을 다루므로 반드시 순차 실행합니다. 각 파일은
독립 Excel 인스턴스에서 처리하며, 성공으로 집계하기 전에 저장 파일을 다시 열어
``exp!T4`` 데이터가 실제로 남아 있는지도 확인합니다.
"""

import gc
import math
import os
import re
import shutil
import threading
import time
from datetime import date, datetime, time as datetime_time
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


# --- 사용자 설정 ---
CONFIG = {
    # RTD 폴더, HX 폴더, _SYS 파일이 들어 있는 실험 원본 최상위 폴더입니다.
    "base_dir": r"E:\HTL\03.PersonalResearch\Supercritical_CO2\04.Raw data\03. Supercritical\20260915_7.771MPa_31.3C_40000",
    # 매크로 원본입니다. 매번 출력 위치로 복사한 뒤 복사본만 수정합니다.
    "xlsm_template_path": r"E:\HTL\03.PersonalResearch\Supercritical_CO2\09.Exp results\YYYYMMDD_n.nnnMPa_nn.nC_Re_nV_Hdn.nn.xlsm",
    # 이 폴더 아래에 날짜/압력/온도/Re별 하위 폴더가 자동 생성됩니다.
    "xlsm_output_dir": r"E:\HTL\03.PersonalResearch\Supercritical_CO2\09.Exp results\03. Supercritical",
    # False이면 Merged.xlsx까지만 만들고 Excel 자동화는 실행하지 않습니다.
    "run_xlsm_automation": True,
    # 안전 기본값입니다. 기존 결과가 있으면 덮어쓰지 않고 건너뜁니다.
    "overwrite_existing_xlsm": False,
    # VBA 팝업 탐지 및 사용자 확인을 위해 기본적으로 Excel 창을 표시합니다.
    "excel_visible": True,
    # Batch_Clear_Data 두 번째 InputBox에 입력할 확인 문자입니다.
    "batch_clear_confirmation": "d",
    # 세 팝업 각각에 적용되는 최대 대기시간입니다.
    "dialog_timeout_seconds": 120,
    # 첫 현장 실행에서 창 클래스와 자식 HWND를 확인할 수 있는 상세 로그입니다.
    "dialog_diagnostic_logging": True,
    # calculator 반환 후 Excel 계산이 끝날 때까지 기다릴 최대 시간입니다.
    "calculation_timeout_seconds": 600,
    "batch_clear_macro": "Batch_Clear_Data",
    "calculator_macro": "calculator",
    # 템플릿 exp 시트에 기록할 실험 조건입니다.
    "chiller_setting_temperature_c": 15.0,  # exp!E5
    "pump_input_hz": 50.0,                  # exp!E6
    "heater_temperature_c": 25.0,           # exp!E8
    "environment_temperature_c": 25.0,      # exp!E9
}


# Excel 워크시트의 물리적 한계입니다. T4 오프셋까지 포함해 사전에 검사합니다.
EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_COLUMNS = 16_384
DESTINATION_START_ROW = 4
DESTINATION_START_COLUMN = 20  # T열

# 파일명 형식이 다르면 값을 추측하지 않습니다. 날짜/조건을 모두 확인한 뒤 처리합니다.
MERGED_FILENAME_PATTERN = re.compile(
    r"^(?P<date>20\d{6})_"
    r"(?P<pressure>[+-]?(?:\d+(?:\.\d+)?|\.\d+))MPa_"
    r"(?P<temperature>[+-]?(?:\d+(?:\.\d+)?|\.\d+))C_"
    r"(?P<reynolds>\d+(?:\.\d+)?)_"
    r"(?P<voltage>[+-]?(?:\d+(?:\.\d+)?|\.\d+))V_"
    r"Hd(?P<hd>[+-]?(?:\d+(?:\.\d+)?|\.\d+))_Merged\.xlsx$",
    re.IGNORECASE,
)


class XlsmSkippedError(Exception):
    """XLSM 후처리를 의도적으로 건너뛸 때 사용하는 예외입니다."""


class _DialogHandlerCancelled(Exception):
    """매크로 호출이 먼저 실패하여 대화상자 감시를 중단했음을 나타냅니다."""


def extract_date_from_str(text):
    """문자열의 YYYYMMDD를 날짜로 바꾸고, 없으면 오늘 날짜를 반환합니다."""
    match = re.search(r"(20\d{2})(\d{2})(\d{2})", text)
    if match:
        return datetime.strptime(match.group(0), "%Y%m%d").date()
    return datetime.now().date()


def parse_time_with_date_injection(time_val, base_date):
    """시간 값에 기준 날짜를 결합해 pandas에서 비교 가능한 값으로 만듭니다."""
    if pd.isna(time_val):
        return pd.NaT
    if isinstance(time_val, datetime):
        if time_val.year == 1900:
            return datetime.combine(base_date, time_val.time())
        return time_val
    if hasattr(time_val, "hour"):
        return datetime.combine(base_date, time_val)
    time_str = str(time_val).strip()
    try:
        return datetime.strptime(time_str, "%Y-%m-%d %H:%M:%S")
    except (TypeError, ValueError):
        pass
    try:
        return datetime.combine(base_date, datetime.strptime(time_str, "%H:%M:%S").time())
    except (TypeError, ValueError):
        return pd.NaT


def process_rtd_dynamic(folders, base_date):
    """기존 RTD 계산 규칙을 유지하면서 폴더별 데이터를 세로로 연결합니다.

    센서 데이터는 T/t/R 세 묶음으로 간주합니다. 센서가 10개보다 적으면 각 묶음
    뒤에 동일한 수의 빈 열을 추가해 기본 31열 구조를 유지하고, 파일 사이에는 빈
    행 하나를 둡니다. 반환되는 시작/종료 시각은 SYS/HX 필터 범위로 사용됩니다.
    """
    all_processed_data = []
    global_start = None
    global_end = None
    folders.sort()

    for folder in folders:
        files = sorted([f for f in os.listdir(folder) if f.endswith(".xlsx") and not f.startswith("~$")])
        for file in files:
            path = os.path.join(folder, file)
            try:
                df_source = pd.read_excel(path, header=0)
                if len(df_source) < 2:
                    continue

                # 첫 열은 시간 정보이므로 제외하고, 나머지를 T/t/R 세 묶음으로 나눕니다.
                total_data_cols = df_source.shape[1] - 1
                num_sensors = total_data_cols // 3
                gap_count = max(0, 10 - num_sensors)
                start_val = df_source.iloc[0, 0]
                end_val = df_source.iloc[1, 0]

                s_dt = parse_time_with_date_injection(start_val, base_date)
                e_dt = parse_time_with_date_injection(end_val, base_date)
                if global_start is None or (s_dt and s_dt < global_start):
                    global_start = s_dt
                if global_end is None or (e_dt and e_dt > global_end):
                    global_end = e_dt

                # 원본 Excel의 31~57행에 해당하는 기존 분석 구간을 그대로 사용합니다.
                slice_end = min(56, len(df_source))
                if slice_end <= 29:
                    continue
                df_sliced = df_source.iloc[29:slice_end].reset_index(drop=True)

                # 병합본 첫 열에는 각 파일의 시작/종료 시각만 남기고 나머지는 비웁니다.
                time_col = pd.Series([np.nan] * len(df_sliced), dtype=object)
                time_col[0] = start_val
                if len(time_col) > 1:
                    time_col[1] = end_val

                idx_t_start = 1
                idx_time_start = 1 + num_sensors
                idx_r_start = 1 + 2 * num_sensors
                t_data = df_sliced.iloc[:, idx_t_start:idx_t_start + num_sensors]
                time_data = df_sliced.iloc[:, idx_time_start:idx_time_start + num_sensors]
                r_data = df_sliced.iloc[:, idx_r_start:idx_r_start + num_sensors]

                if gap_count > 0:
                    gap = pd.DataFrame(np.nan, index=df_sliced.index,
                                       columns=[f"Gap{i}" for i in range(gap_count)])
                else:
                    gap = pd.DataFrame()

                df_combined = pd.concat([time_col, t_data, gap, time_data, gap, r_data, gap], axis=1)
                all_processed_data.append(df_combined)
                blank_row = pd.DataFrame([[np.nan] * df_combined.shape[1]], columns=df_combined.columns)
                all_processed_data.append(blank_row)
            except Exception as exc:
                print(f"    [Error] {file}: {exc}")

    if not all_processed_data:
        return None, None, None
    all_processed_data.pop()
    final_df = pd.concat(all_processed_data, ignore_index=True)
    return final_df, global_start, global_end


def process_filter_data(path, base_date, start_dt, end_dt):
    """SYS 또는 HX 파일에서 RTD 측정 시간에 포함되는 행만 반환합니다."""
    try:
        df = pd.read_excel(path) if path.endswith(".xlsx") else pd.read_csv(path)
    except Exception:
        return None

    time_col = next((c for c in df.columns if "time" in c.lower()), None)
    if not time_col:
        return None
    temp_dates = df[time_col].apply(lambda x: parse_time_with_date_injection(x, base_date))
    if start_dt and end_dt:
        mask = (temp_dates >= start_dt) & (temp_dates <= end_dt)
        return df.loc[mask].copy()
    return pd.DataFrame()


def parse_merged_filename(merged_path):
    """Merged 파일명을 검증하고 XLSM 입력값과 출력 경로 구성 요소를 반환합니다.

    정규식 전체 일치(fullmatch)를 사용하므로 일부 문자열만 우연히 맞는 파일은
    처리하지 않습니다. H/d는 표시 자릿수를 보존하기 위해 숫자가 아닌 문자열로
    유지하며, 날짜 셀 역시 Excel 날짜 객체가 아닌 ``YY.MM.DD`` 문자열을 만듭니다.
    """
    filename = Path(merged_path).name
    match = MERGED_FILENAME_PATTERN.fullmatch(filename)
    if not match:
        raise ValueError(
            "예상 파일명 형식과 일치하지 않습니다: "
            f"{filename!r} (예: YYYYMMDD_7.771MPa_31.3C_40000_8V_Hd1.00_Merged.xlsx)"
        )
    parts = match.groupdict()
    try:
        parsed_date = datetime.strptime(parts["date"], "%Y%m%d")
    except ValueError as exc:
        raise ValueError(f"파일명의 날짜가 유효하지 않습니다: {parts['date']}") from exc

    merged_suffix = "_Merged.xlsx"
    output_filename = filename[:-len(merged_suffix)] + ".xlsm"
    experiment_folder = (
        f"{parts['date']}_{parts['pressure']}MPa_"
        f"{parts['temperature']}C_{parts['reynolds']}"
    )
    return {
        **parts,
        "date_text": parsed_date.strftime("%y.%m.%d"),
        "output_filename": output_filename,
        "experiment_folder": experiment_folder,
    }


def find_actual_data_bounds(worksheet):
    """스타일 전용 셀을 제외하고 실제 값이 있는 마지막 행과 열을 찾습니다.

    openpyxl의 max_row/max_column은 과거에 서식만 적용했던 셀까지 포함할 수 있습니다.
    따라서 실제 생성된 셀 중 값이 있는 셀만 검사해 불필요하게 거대한 COM 범위를
    만드는 문제를 피합니다.
    """
    last_row = 0
    last_column = 0
    # max_row/max_column은 과거 서식 때문에 과장될 수 있으므로 실제 생성 셀의 값을 검사합니다.
    for cell in worksheet._cells.values():
        if cell.value is not None:
            last_row = max(last_row, cell.row)
            last_column = max(last_column, cell.column)
    if last_row == 0 or last_column == 0:
        return None
    return last_row, last_column


def read_merged_values(merged_path):
    """MergedData의 A1부터 실제 마지막 값까지를 메모리의 2차원 tuple로 읽습니다.

    원본 XLSX는 Excel로 열지 않습니다. 중간의 빈 셀/열은 None으로 유지하며, COM에
    한 번에 넘길 수 있도록 직사각형 ``tuple of tuples``를 반환합니다.
    """
    workbook = load_workbook(merged_path, read_only=False, data_only=True)
    try:
        if "MergedData" not in workbook.sheetnames:
            raise ValueError(f"MergedData 시트가 없습니다: {merged_path}")
        worksheet = workbook["MergedData"]
        bounds = find_actual_data_bounds(worksheet)
        if bounds is None:
            raise ValueError(f"MergedData 시트가 비어 있습니다: {merged_path}")

        last_row, last_column = bounds
        if last_row > EXCEL_MAX_ROWS or last_column > EXCEL_MAX_COLUMNS:
            raise ValueError(f"원본 데이터 범위({last_row}행 x {last_column}열)가 Excel 최대 범위를 초과합니다.")
        if DESTINATION_START_ROW + last_row - 1 > EXCEL_MAX_ROWS:
            raise ValueError(f"T4부터 입력할 때 마지막 행이 Excel 한계({EXCEL_MAX_ROWS})를 초과합니다.")
        if DESTINATION_START_COLUMN + last_column - 1 > EXCEL_MAX_COLUMNS:
            raise ValueError(f"T4부터 입력할 때 마지막 열이 Excel 한계({EXCEL_MAX_COLUMNS})를 초과합니다.")

        return tuple(
            tuple(worksheet.cell(row=row, column=column).value for column in range(1, last_column + 1))
            for row in range(1, last_row + 1)
        )
    finally:
        workbook.close()


def normalize_excel_value(value):
    """numpy/pandas 전용 값과 결측값을 Excel COM 호환 Python 값으로 바꿉니다."""
    if value is None:
        return None
    if value is pd.NA or value is pd.NaT:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, (date, datetime)):
        return value
    return value


def _collect_non_empty_cells(values):
    """0-based 상대 좌표와 값을 행 우선 순서로 수집합니다."""
    return [
        (row_index, column_index, value)
        for row_index, row in enumerate(values)
        for column_index, value in enumerate(row)
        if value is not None
    ]


def _select_representative_cells(non_empty_cells):
    """전체 범위에 고르게 퍼진 다섯 셀을 골라 부분 붙여넣기도 탐지합니다."""
    if not non_empty_cells:
        return []
    last_index = len(non_empty_cells) - 1
    indexes = [round(last_index * fraction) for fraction in (0.0, 0.25, 0.5, 0.75, 1.0)]
    return [non_empty_cells[index] for index in indexes]


def _excel_cell_address(row, column):
    """1-based 행/열 번호를 A1 형식 주소로 변환합니다."""
    from openpyxl.utils import get_column_letter

    return f"{get_column_letter(column)}{row}"


def _values_equivalent(source_value, destination_value):
    """COM 형 변환을 고려해 원본 값과 Excel read-back 값이 같은지 판단합니다.

    Excel은 ``20:14:10`` 같은 시간 문자열을 하루에 대한 소수(시간 일련값)로 저장할
    수 있습니다. 이 정상 변환은 허용하되 일반 숫자는 매우 작은 부동소수 오차만
    허용합니다.
    """
    source_value = normalize_excel_value(source_value)
    destination_value = normalize_excel_value(destination_value)
    if source_value is None or destination_value is None:
        return source_value is None and destination_value is None
    if isinstance(source_value, (int, float)) and not isinstance(source_value, bool):
        if isinstance(destination_value, (int, float)) and not isinstance(destination_value, bool):
            return math.isclose(float(source_value), float(destination_value), rel_tol=1e-12, abs_tol=1e-12)
    if isinstance(destination_value, (int, float)) and not isinstance(destination_value, bool):
        source_time = None
        if isinstance(source_value, datetime_time):
            source_time = source_value
        elif isinstance(source_value, str):
            for time_format in ("%H:%M:%S.%f", "%H:%M:%S"):
                try:
                    source_time = datetime.strptime(source_value.strip(), time_format).time()
                    break
                except ValueError:
                    continue
        if source_time is not None:
            seconds = (
                source_time.hour * 3600
                + source_time.minute * 60
                + source_time.second
                + source_time.microsecond / 1_000_000
            )
            return math.isclose(seconds / 86_400, float(destination_value), rel_tol=0.0, abs_tol=1e-9)
    return source_value == destination_value


def _verify_representative_cells(exp_ws, representatives, start_row, start_col, log_prefix):
    """대표 source 셀에 대응하는 Excel 셀을 read-back하고 일치/존재 개수를 반환합니다."""
    matching_count = 0
    non_empty_destination_count = 0
    for source_row_zero, source_col_zero, source_value in representatives:
        destination_row = start_row + source_row_zero
        destination_col = start_col + source_col_zero
        destination_value = exp_ws.Cells(destination_row, destination_col).Value
        source_address = _excel_cell_address(source_row_zero + 1, source_col_zero + 1)
        destination_address = _excel_cell_address(destination_row, destination_col)
        print(
            f"[{log_prefix}] source {source_address}={source_value!r}, "
            f"destination {destination_address}={destination_value!r}"
        )
        if normalize_excel_value(destination_value) is not None:
            non_empty_destination_count += 1
        if _values_equivalent(source_value, destination_value):
            matching_count += 1
    return matching_count, non_empty_destination_count


def _destination_counta(excel, destination):
    """Excel 자체 WorksheetFunction으로 대상 범위의 CountA를 계산합니다."""
    return int(excel.WorksheetFunction.CountA(destination))


def _validate_range_integrity(source_count, destination_count, matching_count,
                              non_empty_destination_count, representative_count, error_message):
    """CountA와 대표 셀을 함께 검사해 부분 손실을 성공으로 오인하지 않게 합니다.

    빈 문자열이나 Excel 형 변환 때문에 CountA가 소폭 달라질 수 있어 0.1%(최소
    5셀)의 차이는 허용합니다. 반면 대표 셀은 모두 존재하고 원본과 일치해야 합니다.
    """
    allowed_count_difference = max(5, math.ceil(source_count * 0.001))
    count_difference = abs(source_count - destination_count)
    if (
        destination_count == 0
        or matching_count != representative_count
        or non_empty_destination_count != representative_count
        or count_difference > allowed_count_difference
    ):
        raise RuntimeError(
            f"{error_message} "
            f"(source CountA={source_count}, destination CountA={destination_count}, "
            f"CountA 차이={count_difference}, 허용 차이={allowed_count_difference}, "
            f"대표 셀 일치={matching_count}/{representative_count})"
        )


def build_output_xlsm_path(merged_path, config=CONFIG):
    """파일명에서 실험 폴더와 최종 XLSM 경로를 구성합니다."""
    parsed = parse_merged_filename(merged_path)
    return Path(config["xlsm_output_dir"]) / parsed["experiment_folder"] / parsed["output_filename"]


def _macro_reference(workbook_name, macro_name):
    """작은따옴표가 든 통합문서 이름도 안전한 Excel 매크로 참조를 만듭니다."""
    if not macro_name or not str(macro_name).strip():
        raise ValueError("매크로 이름이 비어 있습니다.")
    safe_workbook_name = str(workbook_name).replace("'", "''")
    return f"'{safe_workbook_name}'!{str(macro_name).strip()}"


def _normalise_button_text(text):
    """Windows 버튼의 &, (Y) 같은 단축키 표기를 제거합니다."""
    normalised = str(text).strip().casefold().replace("&", "")
    return re.sub(r"\(\s*[a-z]\s*\)", "", normalised).strip()


YES_NAMES = {"예", "yes", "&yes", "예(y)", "예(&y)"}
OK_NAMES = {"확인", "ok", "&ok", "확인(o)", "확인(&o)"}
NORMALISED_YES_NAMES = {_normalise_button_text(name) for name in YES_NAMES}
NORMALISED_OK_NAMES = {_normalise_button_text(name) for name in OK_NAMES}


def enumerate_excel_windows(excel_pid, excel_main_hwnd):
    """같은 Excel PID에 속한 visible top-level 보조 창을 모두 열거합니다.

    Office 버전에 따라 VBA 대화상자 클래스가 #32770 또는 bosa_sdm_XL9 등으로
    달라질 수 있으므로 클래스 이름으로 먼저 거르지 않습니다.
    """
    import win32gui
    import win32process

    windows = []

    def callback(hwnd, _):
        if not win32gui.IsWindowVisible(hwnd):
            return True
        _, pid = win32process.GetWindowThreadProcessId(hwnd)
        if pid != excel_pid or hwnd == excel_main_hwnd:
            return True
        windows.append({
            "hwnd": hwnd,
            "title": win32gui.GetWindowText(hwnd),
            "class_name": win32gui.GetClassName(hwnd),
        })
        return True

    win32gui.EnumWindows(callback, None)
    return windows


def enumerate_child_controls(parent_hwnd):
    """top-level 창 아래의 모든 자식 HWND, 클래스 및 텍스트를 열거합니다."""
    import win32gui

    controls = []

    def callback(hwnd, _):
        controls.append({
            "hwnd": hwnd,
            "class_name": win32gui.GetClassName(hwnd),
            "text": win32gui.GetWindowText(hwnd),
        })
        return True

    win32gui.EnumChildWindows(parent_hwnd, callback, None)
    return controls


def _is_named_button(control, expected_names):
    """열거된 Win32 컨트롤이 원하는 이름의 버튼인지 확인합니다."""
    return (
        control["class_name"].casefold() == "button"
        and _normalise_button_text(control["text"]) in expected_names
    )


def _permission_mismatch_message():
    """Windows UIPI 권한 불일치 해결 안내를 반환합니다."""
    return (
        "Excel과 PyCharm의 실행 권한 수준이 다릅니다.\n"
        "Excel이 관리자 권한으로 실행 중이면 Excel을 종료한 뒤\n"
        "PyCharm도 같은 권한 수준으로 실행하거나, 둘 다 일반 권한으로 실행하세요."
    )


def _send_button_click(button_hwnd):
    """포커스나 키보드에 의존하지 않고 버튼 HWND로 BM_CLICK을 보냅니다."""
    import win32con
    import win32gui

    try:
        win32gui.SendMessage(button_hwnd, win32con.BM_CLICK, 0, 0)
    except Exception as exc:
        print(_permission_mismatch_message())
        raise RuntimeError(f"버튼 HWND 0x{button_hwnd:08X} 클릭 실패: {exc}") from exc


def _send_edit_text(edit_hwnd, text):
    """IME와 포커스에 의존하지 않고 Edit HWND에 WM_SETTEXT를 보냅니다."""
    import win32con
    import win32gui

    try:
        win32gui.SendMessage(edit_hwnd, win32con.WM_SETTEXT, 0, str(text))
    except Exception as exc:
        print(_permission_mismatch_message())
        raise RuntimeError(f"Edit HWND 0x{edit_hwnd:08X} 입력 실패: {exc}") from exc


class BatchClearDialogHandler(threading.Thread):
    """동일 Excel PID의 세 팝업을 자식 컨트롤과 명시적 상태 순서로 처리합니다.

    전역 키보드 입력은 다른 창에 전달될 위험이 있으므로 사용하지 않습니다. Win32
    자식 HWND가 있으면 BM_CLICK/WM_SETTEXT를 직접 보내고, 자식 HWND를 전혀 얻지
    못한 창에 한해서만 UIA backend를 보조 수단으로 사용합니다.
    """

    WAIT_CONFIRM = "WAIT_CONFIRM"
    WAIT_INPUT = "WAIT_INPUT"
    WAIT_COMPLETION = "WAIT_COMPLETION"
    COMPLETED = "COMPLETED"

    def __init__(self, excel_pid, excel_hwnd, confirmation, timeout_seconds,
                 handler_ready, handler_finished, diagnostic_logging=True):
        super().__init__(name=f"BatchClearDialogHandler-{excel_pid}", daemon=True)
        self.excel_pid = excel_pid
        self.excel_hwnd = excel_hwnd
        self.confirmation = str(confirmation)
        self.timeout_seconds = float(timeout_seconds)
        self.handler_ready = handler_ready
        self.handler_finished = handler_finished
        self.diagnostic_logging = bool(diagnostic_logging)
        self.error = None
        self.completed = False
        self.state = self.WAIT_CONFIRM
        self._cancel_event = threading.Event()
        self._logged_snapshots = set()
        self._logged_uia_snapshots = set()

    def cancel(self):
        """매크로 호출 자체가 실패했을 때 감시 스레드를 중단합니다."""
        self._cancel_event.set()

    def _log_window_and_controls(self, window, controls):
        """처음 관찰한 창/컨트롤 구성의 HWND, 클래스 및 텍스트를 진단 출력합니다."""
        signature = (
            self.state,
            window["hwnd"],
            window["title"],
            window["class_name"],
            tuple((item["hwnd"], item["class_name"], item["text"]) for item in controls),
        )
        if not self.diagnostic_logging or signature in self._logged_snapshots:
            return
        self._logged_snapshots.add(signature)
        print(
            f"[Dialog Diagnostic] 상태={self.state}, 부모 HWND=0x{window['hwnd']:08X}, "
            f"부모 제목={window['title']!r}, 부모 클래스={window['class_name']!r}"
        )
        if not controls:
            print("[Dialog Diagnostic] Win32 자식 컨트롤 없음: UIA fallback을 시도합니다.")
        for control in controls:
            print(
                f"[Dialog Diagnostic] 부모 HWND=0x{window['hwnd']:08X}, "
                f"부모 창 제목={window['title']!r}, 부모 클래스={window['class_name']!r}, "
                f"자식 HWND=0x{control['hwnd']:08X}, 자식 클래스={control['class_name']!r}, "
                f"자식 텍스트={control['text']!r}"
            )

    def _perform_win32_action(self, controls):
        """현재 상태와 일치하는 Win32 자식 컨트롤에 직접 메시지를 보냅니다."""
        yes_button = next((c for c in controls if _is_named_button(c, NORMALISED_YES_NAMES)), None)
        ok_button = next((c for c in controls if _is_named_button(c, NORMALISED_OK_NAMES)), None)
        edit = next((c for c in controls if c["class_name"].casefold() == "edit"), None)

        # 각 상태는 해당 단계의 고유 컨트롤을 확인하고 조작에 성공한 뒤에만 전이합니다.
        if self.state == self.WAIT_CONFIRM and yes_button is not None:
            _send_button_click(yes_button["hwnd"])
            self.state = self.WAIT_INPUT
            print("[Batch_Clear_Data] 1/3 삭제 확인 팝업 발견 및 '예' 클릭")
            return True
        if self.state == self.WAIT_INPUT and edit is not None and ok_button is not None:
            _send_edit_text(edit["hwnd"], self.confirmation)
            time.sleep(0.3)
            _send_button_click(ok_button["hwnd"])
            self.state = self.WAIT_COMPLETION
            print(
                f"[Batch_Clear_Data] 2/3 입력 팝업 발견, {self.confirmation!r} 입력 및 '확인' 클릭"
            )
            return True
        if self.state == self.WAIT_COMPLETION and edit is None and ok_button is not None:
            _send_button_click(ok_button["hwnd"])
            self.state = self.COMPLETED
            self.completed = True
            print("[Batch_Clear_Data] 3/3 완료 팝업 발견 및 '확인' 클릭")
            return True
        return False

    def _perform_uia_fallback(self):
        """Win32 자식 HWND가 없을 때만 UIA로 창과 컨트롤을 진단하고 처리합니다."""
        from pywinauto import Desktop

        for window in Desktop(backend="uia").windows(process=self.excel_pid, visible_only=True):
            try:
                handle = int(window.handle or 0)
                if handle == self.excel_hwnd:
                    continue
                descendants = window.descendants()
                uia_items = [
                    (
                        control.window_text(),
                        control.class_name(),
                        control.element_info.control_type,
                        int(control.handle or 0),
                    )
                    for control in descendants
                ]
                signature = (
                    self.state,
                    handle,
                    window.window_text(),
                    window.class_name(),
                    window.element_info.control_type,
                    tuple(uia_items),
                )
                if self.diagnostic_logging and signature not in self._logged_uia_snapshots:
                    self._logged_uia_snapshots.add(signature)
                    print(
                        f"[UIA Fallback] text={window.window_text()!r}, class={window.class_name()!r}, "
                        f"control_type={window.element_info.control_type!r}, handle=0x{handle:08X}"
                    )
                    for text, class_name, control_type, control_handle in uia_items:
                        print(
                            f"[UIA Fallback Child] text={text!r}, class={class_name!r}, "
                            f"control_type={control_type!r}, handle=0x{control_handle:08X}"
                        )

                buttons = [c for c in descendants if c.element_info.control_type == "Button"]
                edits = [c for c in descendants if c.element_info.control_type == "Edit"]
                yes_button = next(
                    (c for c in buttons if _normalise_button_text(c.window_text()) in NORMALISED_YES_NAMES), None
                )
                ok_button = next(
                    (c for c in buttons if _normalise_button_text(c.window_text()) in NORMALISED_OK_NAMES), None
                )

                if self.state == self.WAIT_CONFIRM and yes_button is not None:
                    self._uia_click(yes_button)
                    self.state = self.WAIT_INPUT
                    print("[Batch_Clear_Data] 1/3 삭제 확인 팝업 발견 및 '예' 클릭 (UIA)")
                    return True
                if self.state == self.WAIT_INPUT and edits and ok_button is not None:
                    self._uia_set_text(edits[0], self.confirmation)
                    time.sleep(0.3)
                    self._uia_click(ok_button)
                    self.state = self.WAIT_COMPLETION
                    print(
                        f"[Batch_Clear_Data] 2/3 입력 팝업 발견, {self.confirmation!r} "
                        "입력 및 '확인' 클릭 (UIA)"
                    )
                    return True
                if self.state == self.WAIT_COMPLETION and not edits and ok_button is not None:
                    self._uia_click(ok_button)
                    self.state = self.COMPLETED
                    self.completed = True
                    print("[Batch_Clear_Data] 3/3 완료 팝업 발견 및 '확인' 클릭 (UIA)")
                    return True
            except Exception as exc:
                print(f"[UIA Fallback] 창 검사/조작 실패: {exc}")
                if "access" in str(exc).casefold() or "denied" in str(exc).casefold():
                    print(_permission_mismatch_message())
        return False

    @staticmethod
    def _uia_click(control):
        """UIA 컨트롤을 전역 키 입력 없이 invoke/click합니다."""
        try:
            control.invoke()
        except Exception:
            control.click()

    @staticmethod
    def _uia_set_text(control, text):
        """UIA Edit 컨트롤 값을 전역 키 입력 없이 설정합니다."""
        control.set_edit_text(str(text))

    def _dismiss_process_windows(self):
        """오류 시 동일 PID의 보조 창만 닫아 블로킹 매크로 호출을 해제합니다."""
        import win32con
        import win32gui

        windows = enumerate_excel_windows(self.excel_pid, self.excel_hwnd)
        for window in windows:
            try:
                win32gui.PostMessage(window["hwnd"], win32con.WM_CLOSE, 0, 0)
            except Exception:
                pass
        if not windows:
            try:
                win32gui.PostMessage(self.excel_hwnd, win32con.WM_CLOSE, 0, 0)
            except Exception:
                pass

    def run(self):
        """WAIT_CONFIRM부터 COMPLETED까지 단계별 timeout을 적용해 감시합니다."""
        try:
            import win32gui  # 감시 시작 전에 pywin32 사용 가능 여부를 확인합니다.
            import win32process

            _, actual_pid = win32process.GetWindowThreadProcessId(self.excel_hwnd)
            if actual_pid != self.excel_pid:
                raise RuntimeError(
                    f"Excel HWND의 PID가 변경되었습니다: expected={self.excel_pid}, actual={actual_pid}"
                )
            self.handler_ready.set()
            print("[Dialog Handler] 감시 시작")

            state_deadline = time.monotonic() + self.timeout_seconds
            next_empty_log = 0.0
            while self.state != self.COMPLETED:
                if self._cancel_event.is_set():
                    raise _DialogHandlerCancelled()
                if time.monotonic() >= state_deadline:
                    raise TimeoutError(
                        f"Batch_Clear_Data 팝업 단계 {self.state}를 "
                        f"{self.timeout_seconds:g}초 안에 처리하지 못했습니다."
                    )

                # Python이 만든 Excel PID만 열거하므로 사용자가 연 다른 Excel은 건드리지 않습니다.
                windows = enumerate_excel_windows(self.excel_pid, self.excel_hwnd)
                now = time.monotonic()
                if not windows:
                    if now >= next_empty_log:
                        print("[Dialog Handler] 현재 PID에서 별도 visible top-level window를 찾지 못했습니다.")
                        next_empty_log = now + 1.0
                    time.sleep(0.1)
                    continue

                action_performed = False
                fallback_needed = False
                for window in windows:
                    controls = enumerate_child_controls(window["hwnd"])
                    self._log_window_and_controls(window, controls)
                    if not controls:
                        fallback_needed = True
                        continue
                    if self._perform_win32_action(controls):
                        action_performed = True
                        break

                if not action_performed and fallback_needed:
                    action_performed = self._perform_uia_fallback()
                if action_performed:
                    state_deadline = time.monotonic() + self.timeout_seconds
                else:
                    time.sleep(0.1)
        except _DialogHandlerCancelled:
            return
        except BaseException as exc:
            self.error = exc
            self.handler_ready.set()
            self._dismiss_process_windows()
        finally:
            self.handler_finished.set()


def run_batch_clear_with_dialog_handler(excel, workbook, excel_pid, config=CONFIG):
    """팝업 감시 준비를 확인한 뒤 블로킹 방식으로 Batch_Clear_Data를 실행합니다.

    VBA MsgBox/InputBox가 떠 있는 동안 Application.Run은 반환하지 않을 수 있습니다.
    따라서 감시 스레드가 준비됐다는 Event를 받은 뒤에만 매크로를 호출합니다.
    """
    import win32con
    import win32gui

    excel_main_hwnd = int(excel.Hwnd)
    macro_ref = _macro_reference(workbook.Name, config["batch_clear_macro"])
    workbook.Activate()
    excel.Visible = True
    try:
        excel.UserControl = True
    except Exception as exc:
        print(f"[Dialog Handler] Excel UserControl 설정 참고: {exc}")
    try:
        win32gui.ShowWindow(excel_main_hwnd, win32con.SW_RESTORE)
        win32gui.SetForegroundWindow(excel_main_hwnd)
    except Exception as exc:
        print(f"[Dialog Handler] Excel 창 활성화 참고: {exc}")

    # ready는 매크로 시작 전 경쟁 조건을 막고, finished는 스레드 예외 회수를 보장합니다.
    handler_ready = threading.Event()
    handler_finished = threading.Event()
    handler = BatchClearDialogHandler(
        excel_pid=excel_pid,
        excel_hwnd=excel_main_hwnd,
        confirmation=config["batch_clear_confirmation"],
        timeout_seconds=config["dialog_timeout_seconds"],
        handler_ready=handler_ready,
        handler_finished=handler_finished,
        diagnostic_logging=config.get("dialog_diagnostic_logging", True),
    )
    print(f"[Dialog Handler] Excel PID: {excel_pid}")
    print(f"[Dialog Handler] Excel HWND: {excel_main_hwnd:08X}")
    print(f"[Dialog Handler] Workbook: {workbook.Name}")
    print(f"[Dialog Handler] 팝업 처리 thread 시작 여부: 시작 요청")
    handler.start()
    if not handler_ready.wait(timeout=10.0):
        handler.cancel()
        handler.join(timeout=5.0)
        raise RuntimeError("팝업 처리 스레드가 준비되지 않았습니다.")
    if handler.error is not None:
        handler.join(timeout=5.0)
        raise RuntimeError(f"Batch_Clear_Data 팝업 처리기 초기화 실패: {handler.error}") from handler.error
    print("[Dialog Handler] 팝업 처리 thread 시작 여부: 준비 완료")
    time.sleep(0.5)
    print(f"[Macro] {macro_ref} 실행")
    macro_error = None
    try:
        excel.Application.Run(macro_ref)
    except BaseException as exc:
        macro_error = exc
        handler.cancel()

    handler_finished.wait(timeout=2.0 if macro_error is None else 5.0)
    handler.join(timeout=0.5)
    if handler.is_alive():
        handler.cancel()
        handler.join(timeout=5.0)
    if handler.is_alive():
        raise TimeoutError("Batch_Clear_Data 팝업 처리 스레드가 종료되지 않았습니다.")
    if handler.error is not None:
        raise RuntimeError(f"Batch_Clear_Data 팝업 처리 실패: {handler.error}") from handler.error
    if macro_error is not None:
        if isinstance(macro_error, KeyboardInterrupt):
            raise macro_error
        raise RuntimeError(f"매크로 실행 실패 ({macro_ref}): {macro_error}") from macro_error
    if not handler.completed:
        raise RuntimeError("Batch_Clear_Data가 반환됐지만 세 개의 팝업 처리가 완료되지 않았습니다.")


def wait_for_excel_calculation(excel, timeout_seconds):
    """비동기 쿼리를 정리한 뒤 xlDone까지 기다리되 무한 대기는 방지합니다."""
    try:
        excel.CalculateUntilAsyncQueriesDone()
    except Exception as exc:
        print(f"    [계산 참고] 비동기 쿼리 완료 확인을 사용할 수 없습니다: {exc}")
    deadline = time.monotonic() + float(timeout_seconds)
    while time.monotonic() < deadline:
        if int(excel.CalculationState) == 0:
            return
        time.sleep(0.25)
    raise TimeoutError(f"Excel 계산이 {float(timeout_seconds):g}초 안에 완료되지 않았습니다.")


def _excel_number(number_text):
    """파일명의 숫자 문자열을 Excel용 int 또는 float로 변환합니다."""
    return int(number_text) if "." not in number_text else float(number_text)


def create_and_process_xlsm(merged_path, config=CONFIG):
    """Merged.xlsx 하나를 XLSM으로 후처리하고 저장 결과까지 검증합니다.

    처리 순서는 템플릿 복사 → Batch_Clear_Data → exp 입력 → calculator → 저장 →
    read-only 재열기입니다. 어느 단계에서든 실패하면 성공을 반환하지 않으며, 조사할
    수 있도록 복사된 XLSM은 삭제하지 않습니다. Excel/workbook 정리는 finally에서
    다시 시도해 Python이 만든 Excel 프로세스가 남을 가능성을 줄입니다.
    """
    merged_path = Path(merged_path)
    try:
        parsed = parse_merged_filename(merged_path)
    except ValueError as exc:
        raise XlsmSkippedError(str(exc)) from exc

    template_path = Path(config["xlsm_template_path"])
    if not template_path.is_file():
        raise FileNotFoundError(f"XLSM 템플릿 파일이 없습니다: {template_path}")
    if not merged_path.is_file():
        raise FileNotFoundError(f"Merged.xlsx 파일이 없습니다: {merged_path}")

    # generator나 pandas/numpy 전용 값이 SAFEARRAY에 섞이지 않도록 미리 완전히 변환합니다.
    values = tuple(tuple(row) for row in read_merged_values(merged_path))
    excel_values = tuple(
        tuple(normalize_excel_value(value) for value in row)
        for row in values
    )
    row_count = len(values)
    column_count = len(values[0])
    non_empty_cells = _collect_non_empty_cells(excel_values)
    if not non_empty_cells:
        raise ValueError(f"Merged.xlsx에 COM으로 전달할 실제 데이터가 없습니다: {merged_path}")
    representatives = _select_representative_cells(non_empty_cells)
    first_row, first_col, first_value = non_empty_cells[0]
    last_row, last_col, last_value = non_empty_cells[-1]
    print(f"[Source 검증] {row_count}행 x {column_count}열")
    print(f"[Source 검증] 비어 있지 않은 셀: {len(non_empty_cells)}")
    print(f"[Source 검증] 첫 값: {_excel_cell_address(first_row + 1, first_col + 1)} = {first_value!r}")
    print(f"[Source 검증] 마지막 값: {_excel_cell_address(last_row + 1, last_col + 1)} = {last_value!r}")

    output_path = build_output_xlsm_path(merged_path, config)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if output_path.exists() and not config["overwrite_existing_xlsm"]:
        raise XlsmSkippedError(f"기존 결과 보호를 위해 건너뜁니다: {output_path}")

    import pythoncom
    import win32process
    from win32com.client import DispatchEx

    # 원본 템플릿의 VBA를 보존하기 위해 파일 자체를 먼저 복사하고 복사본만 COM으로 엽니다.
    shutil.copy2(template_path, output_path)
    print(f"  [1/6] 템플릿 복사 완료: {output_path}")
    excel = None
    workbook = None
    verify_workbook = None
    exp_ws = None
    verify_ws = None
    destination = None
    destination_after = None
    destination_verify = None
    pythoncom.CoInitialize()
    try:
        # DispatchEx는 사용자가 이미 열어 둔 Excel과 분리된 새 인스턴스를 만듭니다.
        excel = DispatchEx("Excel.Application")
        excel.Visible = bool(config["excel_visible"])
        excel.DisplayAlerts = False
        workbook = excel.Workbooks.Open(str(output_path), UpdateLinks=0, ReadOnly=False)
        if int(excel.Workbooks.Count) != 1:
            raise RuntimeError(f"독립 Excel 인스턴스에 예상치 못한 통합문서가 열렸습니다 (Workbooks.Count={excel.Workbooks.Count}).")
        excel_pid = win32process.GetWindowThreadProcessId(int(excel.Hwnd))[1]

        run_batch_clear_with_dialog_handler(excel, workbook, excel_pid, config)
        print("  [2/6] Batch_Clear_Data 실행 및 팝업 처리 완료")
        try:
            exp_ws = workbook.Worksheets("exp")
        except Exception as exc:
            raise ValueError(f"템플릿 XLSM에 exp 시트가 없습니다: {output_path}") from exc

        exp_ws.Range("B1").Value = parsed["hd"]
        exp_ws.Range("E3").Value = parsed["date_text"]
        exp_ws.Range("E4").Value = _excel_number(parsed["voltage"])
        exp_ws.Range("E5").Value = config["chiller_setting_temperature_c"]
        exp_ws.Range("E6").Value = config["pump_input_hz"]
        exp_ws.Range("E8").Value = config["heater_temperature_c"]
        exp_ws.Range("E9").Value = config["environment_temperature_c"]
        print("  [3/6] exp 기본값 입력 완료")
        print(f"        B1 H/d: {parsed['hd']}")
        print(f"        E3 날짜: {parsed['date_text']}")
        print(f"        E4 전압: {parsed['voltage']}")
        print(f"        E5 Chiller 설정온도: {config['chiller_setting_temperature_c']} °C")
        print(f"        E6 Pump 입력: {config['pump_input_hz']} Hz")
        print(f"        E8 Heater 온도: {config['heater_temperature_c']} °C")
        print(f"        E9 환경온도: {config['environment_temperature_c']} °C")

        # source 좌표는 0-based, Excel Cells/Range 좌표는 1-based입니다.
        start_row = 4
        start_col = 20
        end_row = start_row + row_count - 1
        end_col = start_col + column_count - 1
        workbook.Activate()
        exp_ws.Activate()
        destination = exp_ws.Range(
            exp_ws.Cells(start_row, start_col),
            exp_ws.Cells(end_row, end_col),
        )
        # 셀별 반복은 매우 느리므로 직사각형 범위에 2차원 배열을 한 번만 할당합니다.
        destination.Value = excel_values
        print(f"  [4/6] MergedData {row_count}행 x {column_count}열 범위 할당")

        immediate_matches, immediate_non_empty = _verify_representative_cells(
            exp_ws, representatives, start_row, start_col, "Paste 검증 전"
        )
        source_count = len(non_empty_cells)
        destination_count = _destination_counta(excel, destination)
        print(f"[Paste 검증 전] source CountA: {source_count}")
        print(f"[Paste 검증 전] destination CountA: {destination_count}")
        _validate_range_integrity(
            source_count,
            destination_count,
            immediate_matches,
            immediate_non_empty,
            len(representatives),
            "MergedData 범위 할당문은 반환됐지만 exp!T4에서 실제 값을 확인하지 못했습니다.",
        )
        print("      입력 직후 검증 완료")

        # calculator 실행 전에 한 번 저장해 입력 데이터가 저장 단계에서도 유지되는지 확인합니다.
        workbook.Save()
        saved_matches, saved_non_empty = _verify_representative_cells(
            exp_ws, representatives, start_row, start_col, "저장 검증 전"
        )
        saved_count = _destination_counta(excel, destination)
        _validate_range_integrity(
            source_count,
            saved_count,
            saved_matches,
            saved_non_empty,
            len(representatives),
            "임시 저장 직후 exp!T4 데이터가 유지되지 않았습니다.",
        )
        print("[저장 검증 전] exp!T4 데이터 유지 확인")

        destination = exp_ws.Range(
            exp_ws.Cells(start_row, start_col),
            exp_ws.Cells(end_row, end_col),
        )
        count_before_calculator = _destination_counta(excel, destination)
        print(f"[Calculator 전] destination CountA: {count_before_calculator}")

        # calculator는 ActiveSheet에 의존하므로 대상 workbook과 exp 시트를 명시적으로 활성화합니다.
        workbook.Activate()
        exp_ws.Activate()
        calculator_ref = _macro_reference(workbook.Name, config["calculator_macro"])
        active_workbook_before = excel.ActiveWorkbook.Name if excel.ActiveWorkbook is not None else "(없음)"
        active_worksheet_before = excel.ActiveSheet.Name if excel.ActiveSheet is not None else "(없음)"
        print(f"[Calculator 진단] 실제 호출 매크로: {calculator_ref}")
        print(f"[Calculator 진단] 실행 전 active workbook: {active_workbook_before}")
        print(f"[Calculator 진단] 실행 전 active worksheet: {active_worksheet_before}")
        try:
            excel.Application.Run(calculator_ref)
        except KeyboardInterrupt:
            raise
        except BaseException as exc:
            raise RuntimeError(f"매크로 실행 실패 ({calculator_ref}): {exc}") from exc
        wait_for_excel_calculation(excel, config["calculation_timeout_seconds"])
        active_workbook_after = excel.ActiveWorkbook.Name if excel.ActiveWorkbook is not None else "(없음)"
        active_worksheet_after = excel.ActiveSheet.Name if excel.ActiveSheet is not None else "(없음)"
        print(f"[Calculator 진단] 실행 후 active workbook: {active_workbook_after}")
        print(f"[Calculator 진단] 실행 후 active worksheet: {active_worksheet_after}")

        exp_ws = workbook.Worksheets("exp")
        destination_after = exp_ws.Range(
            exp_ws.Cells(start_row, start_col),
            exp_ws.Cells(end_row, end_col),
        )
        count_after_calculator = _destination_counta(excel, destination_after)
        print(f"[Calculator 후] destination CountA: {count_after_calculator}")
        after_matches, after_non_empty = _verify_representative_cells(
            exp_ws, representatives, start_row, start_col, "Calculator 후"
        )
        try:
            _validate_range_integrity(
                source_count,
                count_after_calculator,
                after_matches,
                after_non_empty,
                len(representatives),
                "calculator 실행 후 exp!T4 데이터가 삭제되었습니다.",
            )
        except RuntimeError:
            print(f"[Calculator 진단] 실제 호출 매크로: {calculator_ref}")
            print(f"[Calculator 진단] 실행 전 active workbook: {active_workbook_before}")
            print(f"[Calculator 진단] 실행 전 active worksheet: {active_worksheet_before}")
            print(f"[Calculator 진단] 실행 후 active workbook: {active_workbook_after}")
            print(f"[Calculator 진단] 실행 후 active worksheet: {active_worksheet_after}")
            raise
        print("  [5/6] calculator 실행 및 계산 완료")

        # calculator 결과와 입력 데이터가 모두 유지된 상태를 최종 저장합니다.
        workbook.Save()
        final_save_matches, final_save_non_empty = _verify_representative_cells(
            exp_ws, representatives, start_row, start_col, "최종 저장 직후"
        )
        final_save_count = _destination_counta(excel, destination_after)
        _validate_range_integrity(
            source_count,
            final_save_count,
            final_save_matches,
            final_save_non_empty,
            len(representatives),
            "최종 저장 직후 exp!T4 데이터가 유지되지 않았습니다.",
        )

        destination = None
        exp_ws = None
        workbook.Close(SaveChanges=True)
        workbook = None

        # 메모리상의 COM 값이 아니라 디스크에 기록된 결과를 확인하기 위해 같은 Excel에서 재오픈합니다.
        verify_workbook = excel.Workbooks.Open(str(output_path), UpdateLinks=0, ReadOnly=True)
        try:
            verify_ws = verify_workbook.Worksheets("exp")
        except Exception as exc:
            raise ValueError(f"재열기한 XLSM에 exp 시트가 없습니다: {output_path}") from exc
        destination_verify = verify_ws.Range(
            verify_ws.Cells(start_row, start_col),
            verify_ws.Cells(end_row, end_col),
        )
        final_count = _destination_counta(excel, destination_verify)
        print(f"[재열기 검증] destination CountA: {final_count}")
        reopened_matches, reopened_non_empty = _verify_representative_cells(
            verify_ws, representatives, start_row, start_col, "재열기 검증"
        )
        _validate_range_integrity(
            source_count,
            final_count,
            reopened_matches,
            reopened_non_empty,
            len(representatives),
            "저장된 xlsm 파일을 다시 열었을 때 exp!T4 데이터가 확인되지 않습니다.",
        )
        print("[재열기 검증] 대표 셀 값 유지 확인")
        print(f"      calculator 전 CountA: {count_before_calculator}")
        print(f"      calculator 후 CountA: {count_after_calculator}")
        print(f"      최종 재열기 CountA: {final_count}")

        destination_verify = None
        verify_ws = None
        verify_workbook.Close(SaveChanges=False)
        verify_workbook = None
        excel.Quit()
        excel = None
        print("  [6/6] 저장 및 Excel 종료 완료")
        return output_path
    except BaseException:
        print(f"  [미완성 XLSM 보존] 확인이 필요한 파일: {output_path}")
        raise
    finally:
        destination = None
        destination_after = None
        destination_verify = None
        exp_ws = None
        verify_ws = None
        if verify_workbook is not None:
            try:
                verify_workbook.Close(SaveChanges=False)
            except Exception as exc:
                print(f"  [정리 경고] 검증 workbook 닫기 실패: {exc}")
        verify_workbook = None
        if workbook is not None:
            try:
                workbook.Close(SaveChanges=False)
            except Exception as exc:
                print(f"  [정리 경고] workbook 닫기 실패: {exc}")
        workbook = None
        if excel is not None:
            try:
                excel.Quit()
            except Exception as exc:
                print(f"  [정리 경고] Python이 생성한 Excel 인스턴스 종료 실패: {exc}")
        excel = None
        gc.collect()
        pythoncom.CoUninitialize()


def main():
    """기존 병합을 수행하고 저장에 성공한 Merged.xlsx만 순차 후처리합니다."""
    base_dir = CONFIG["base_dir"]
    rtd_dir = os.path.join(base_dir, "RTD")
    hx_dir = os.path.join(base_dir, "HX")
    created_merged_files = []

    sys_files = [f for f in os.listdir(base_dir) if f.startswith("_SYS") and f.endswith(".xlsx")]
    sys_path = os.path.join(base_dir, sys_files[0]) if sys_files else None
    rtd_groups = {}
    if os.path.exists(rtd_dir):
        for directory in os.listdir(rtd_dir):
            if "-r" in directory and directory.endswith("mm"):
                match = re.search(r"_RTD-(.+)-r\d+mm", directory)
                if match:
                    key = match.group(1)
                    if key not in rtd_groups:
                        rtd_groups[key] = []
                    rtd_groups[key].append(os.path.join(rtd_dir, directory))

    print(f"--- 총 {len(rtd_groups)}개의 실험 세트 처리 시작 ---\n")
    for group_key, folders in rtd_groups.items():
        print(f">>> 처리 중: {group_key}")
        sample_file = os.listdir(folders[0])[0]
        base_date = extract_date_from_str(sample_file)

        # A. RTD 처리(기존 동적 Gap 계산)
        rtd_df, start_dt, end_dt = process_rtd_dynamic(folders, base_date)
        if rtd_df is None:
            continue
        print(f"    시간 범위: {start_dt.time()} ~ {end_dt.time()}")

        # B. System 처리
        sys_df = process_filter_data(sys_path, base_date, start_dt, end_dt)

        # C. Heater 처리
        heater_df = None
        if os.path.exists(hx_dir):
            hx_candidates = [f for f in os.listdir(hx_dir) if group_key in f and f.endswith(".xlsx")]
            if hx_candidates:
                heater_df = process_filter_data(os.path.join(hx_dir, hx_candidates[0]), base_date, start_dt, end_dt)

        # D. 기존 Merged.xlsx 저장
        output_filename = f"{base_date.strftime('%Y%m%d')}_{group_key}_Merged.xlsx"
        output_path = os.path.join(base_dir, output_filename)
        try:
            with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
                rtd_df.to_excel(writer, sheet_name="MergedData", startrow=0, startcol=0,
                                index=False, header=False)
                sys_start_col = rtd_df.shape[1] + 1
                if sys_df is not None and not sys_df.empty:
                    sys_df.to_excel(writer, sheet_name="MergedData", startrow=0, startcol=sys_start_col,
                                    index=False, header=False)
                sys_width = len(sys_df.columns) if (sys_df is not None and not sys_df.empty) else 0
                heater_start_col = sys_start_col + sys_width + 1
                if heater_df is not None and not heater_df.empty:
                    heater_df.to_excel(writer, sheet_name="MergedData", startrow=0, startcol=heater_start_col,
                                       index=False, header=False)

            # ExcelWriter가 예외 없이 닫힌 파일만 XLSM 후처리 대상으로 등록합니다.
            created_merged_files.append(Path(output_path))
            print(f"    [생성 완료] {output_filename}")
        except Exception as exc:
            print(f"    [저장 실패] {exc}")
        print("")

    xlsm_success_count = 0
    xlsm_failure_count = 0
    skipped_count = 0
    if CONFIG["run_xlsm_automation"]:
        total = len(created_merged_files)
        # Excel COM은 병렬 실행하지 않습니다. 한 파일의 Quit까지 끝난 뒤 다음 파일로 갑니다.
        for index, merged_path in enumerate(created_merged_files, start=1):
            print(f"[XLSM {index}/{total}] 처리 시작: {merged_path.name}")
            try:
                output_path = create_and_process_xlsm(merged_path, CONFIG)
                xlsm_success_count += 1
                print(f"[XLSM 성공] {output_path}\n")
            except XlsmSkippedError as exc:
                skipped_count += 1
                print(f"[XLSM 건너뜀] {exc}\n")
            except KeyboardInterrupt:
                xlsm_failure_count += 1
                print(f"[XLSM 중단] 사용자 요청으로 중단했습니다: {merged_path}")
                break
            except Exception as exc:
                xlsm_failure_count += 1
                print(f"[XLSM 실패] {merged_path}: {exc}")
                print(
                    "  매크로 보안 관련 오류라면 Excel 신뢰 센터에서 템플릿/출력 폴더를 "
                    "'신뢰할 수 있는 위치'로 등록하세요. 프로그램은 보안 정책을 변경하지 않습니다.\n"
                )
    else:
        print("[XLSM] run_xlsm_automation=False: Merged.xlsx 생성까지만 수행합니다.\n")

    print("--- 처리 요약 ---")
    print(f"Merged.xlsx 생성 성공 수: {len(created_merged_files)}")
    print(f"xlsm 처리 성공 수: {xlsm_success_count}")
    print(f"xlsm 처리 실패 수: {xlsm_failure_count}")
    print(f"건너뛴 파일 수: {skipped_count}")


if __name__ == "__main__":
    main()
